"""Optional document intelligence (enabled with RAG_ENABLED=true).

PDF/DOCX -> text -> overlapping chunks -> Gemini embeddings -> stored in MySQL (float32 blobs).
Retrieval is owner-scoped cosine similarity with numpy. The VectorStore surface is small
(`ingest`, `search`, `delete`), so FAISS or ChromaDB can replace the numpy scan for large corpora.
"""
from __future__ import annotations

import io
import logging
import re

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from .config import Settings
from .llm import LLMClient, LLMError
from .models import Document, DocumentChunk

log = logging.getLogger("omniagent.rag")

ALLOWED = {
    ".pdf": ("application/pdf", b"%PDF-"),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK\x03\x04"),
}


class DocumentError(Exception):
    """User-facing document problem."""


def validate_upload(filename: str, data: bytes, settings: Settings) -> tuple[str, str]:
    name = re.sub(r"[^\w.\- ]", "_", (filename or "document").split("/")[-1].split("\\")[-1])[:200]
    ext = "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in ALLOWED:
        raise DocumentError("Only PDF and DOCX files are supported.")
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise DocumentError(f"File is larger than {settings.max_upload_mb} MB.")
    content_type, magic = ALLOWED[ext]
    if not data.startswith(magic):
        raise DocumentError("The file content does not match its extension.")
    return name, content_type


def extract_text(name: str, data: bytes) -> str:
    try:
        if name.lower().endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise DocumentError("Encrypted PDFs are not supported.")
            pages = [(p.extract_text() or "") for p in reader.pages[:400]]
            return "\n\n".join(pages)
        from docx import Document as Docx

        doc = Docx(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(parts)
    except DocumentError:
        raise
    except Exception:
        log.exception("text extraction failed")
        raise DocumentError("Could not read this file. It may be corrupted.") from None


def chunk_text(text: str, size: int, overlap: int) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):  # prefer to break on a paragraph / sentence boundary
            window = text[start:end]
            cut = max(window.rfind("\n\n"), window.rfind(". "), window.rfind("\n"))
            if cut > size * 0.5:
                end = start + cut + 1
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


class RagService:
    def __init__(self, settings: Settings, llm: LLMClient) -> None:
        self.settings = settings
        self.llm = llm

    def ingest(self, db: Session, user_id: int, domain_id: int | None, filename: str, data: bytes) -> Document:
        import numpy as np

        name, content_type = validate_upload(filename, data, self.settings)
        text = extract_text(name, data)
        chunks = chunk_text(text, self.settings.rag_chunk_chars, self.settings.rag_chunk_overlap)
        if not chunks:
            raise DocumentError("No extractable text found (scanned documents are not supported).")
        if len(chunks) > 1500:
            raise DocumentError("This document is too large to index.")
        try:
            vectors = self.llm.embed(chunks, "RETRIEVAL_DOCUMENT")
        except LLMError as e:
            raise DocumentError(str(e)) from None
        doc = Document(
            user_id=user_id, domain_id=domain_id, filename=name, content_type=content_type,
            size_bytes=len(data), chunk_count=len(chunks),
        )
        db.add(doc)
        db.flush()
        for i, (c, v) in enumerate(zip(chunks, vectors)):
            vec = np.asarray(v, dtype="float32")
            vec = vec / (np.linalg.norm(vec) or 1.0)
            db.add(DocumentChunk(document_id=doc.id, user_id=user_id, chunk_index=i, text=c, embedding=vec.tobytes()))
        db.commit()
        return doc

    def search(self, db: Session, user_id: int, query: str, top_k: int | None = None) -> list[dict]:
        import numpy as np

        top_k = max(1, min(top_k or self.settings.rag_top_k, 10))
        rows = db.execute(
            select(DocumentChunk, Document.filename)
            .join(Document, Document.id == DocumentChunk.document_id)
            .where(DocumentChunk.user_id == user_id)  # ownership enforced here
        ).all()
        if not rows:
            return []
        q = np.asarray(self.llm.embed([query], "RETRIEVAL_QUERY")[0], dtype="float32")
        q = q / (np.linalg.norm(q) or 1.0)
        mat = np.vstack([np.frombuffer(r[0].embedding, dtype="float32") for r in rows])
        scores = mat @ q
        best = np.argsort(-scores)[:top_k]
        return [
            {
                "document": rows[i][1],
                "chunk": rows[i][0].chunk_index + 1,
                "score": round(float(scores[i]), 3),
                "text": rows[i][0].text,
            }
            for i in best
            if scores[i] > 0.25
        ]

    def delete(self, db: Session, user_id: int, doc_id: str) -> bool:
        doc = db.get(Document, doc_id)
        if doc is None or doc.user_id != user_id:
            return False
        db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc_id))
        db.delete(doc)
        db.commit()
        return True
